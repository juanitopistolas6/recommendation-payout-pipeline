"""Registro de documentos: vincula cada archivo de data/claims/<tracking_id>/ con su claim y lo deja listo
para Gemini. No llama a ningún LLM.

Por archivo:
    - Hash SHA-256. Si ya hay un ClaimDocument con ese (claim, file_hash), se salta.
    - docx, doc, rtf (.rmd) y xlsx -> PDF con LibreOffice, en _converted/<tracking_id>/<ruta>.<ext>.pdf
      (claims/1/Lease.docx -> _converted/1/Lease.docx.pdf; ver claims/paths.py).
    - HEIC -> JPG con pillow-heif, con la misma regla (_converted/1/IMG_1.HEIC.jpg).
    - PDF (original o convertido): páginas, texto de las 2 primeras, texto completo y has_scanned_pages.
    - Imágenes: page_count = 1, sin texto.
    - Otros tipos (zip, txt, msg, ...): no se registran; se reportan como no soportados.

Un archivo corrupto o que no se puede convertir no detiene el proceso: queda en report.errors.
"""

from __future__ import annotations

import hashlib
import logging
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from claims import paths
from claims.models import Claim, ClaimDocument

logger = logging.getLogger(__name__)

SCANNED_PAGE_MIN_CHARS = 100
SOFFICE_TIMEOUT = 180  # segundos por archivo

MIME_BY_EXTENSION = {
    ".pdf": "application/pdf",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".jfif": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".heic": "image/heic",
    ".heif": "image/heif",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".rtf": "application/rtf",
    ".rmd": "application/rtf",  # en estos datos los .rmd son RTF (se verifica el contenido)
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".jfif", ".png", ".gif", ".webp"}
HEIC_EXTENSIONS = {".heic", ".heif"}
# Extensión con la que LibreOffice recibe el archivo (los .rmd son RTF con otra extensión)
OFFICE_EXTENSIONS = {".docx": ".docx", ".doc": ".doc", ".rtf": ".rtf", ".rmd": ".rtf", ".xlsx": ".xlsx"}

IGNORED_NAMES = {".ds_store", "thumbs.db", "desktop.ini"}


@dataclass
class RegistrationReport:
    tracking_id: str
    folder_found: bool = True
    registered: int = 0
    skipped: int = 0
    converted: int = 0
    unsupported: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def claim_folder(tracking_id: str) -> Path:
    return paths.claims_dir() / tracking_id


def file_kind(path: Path) -> str | None:
    """pdf | image | heic | office, o None si no está soportado."""
    ext = path.suffix.lower()
    if ext in PDF_EXTENSIONS:
        return "pdf"
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in HEIC_EXTENSIONS:
        return "heic"
    if ext in OFFICE_EXTENSIONS:
        if ext == ".rmd":
            with path.open("rb") as fh:
                if fh.read(5) != b"{\\rtf":
                    return None  # un .rmd que no es RTF (p. ej. R Markdown de verdad)
        return "office"
    return None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def register_by_tracking_id(tracking_id: str) -> RegistrationReport:
    """Para procesos en paralelo (register_documents --workers N): recibe solo el ID, que se puede pasar
    entre procesos, y abre su propia conexión a la BD."""
    return register_claim_documents(Claim.objects.get(pk=tracking_id))


def register_claim_documents(claim: Claim) -> RegistrationReport:
    report = RegistrationReport(claim.pk)
    folder = claim_folder(claim.pk)
    if not folder.is_dir():
        report.folder_found = False
        return report

    known = set(claim.documents.values_list("file_hash", flat=True))
    for path in sorted(p for p in folder.rglob("*") if p.is_file()):
        name = path.name.lower()
        if name in IGNORED_NAMES or name.startswith("~$"):  # archivos de sistema y temporales de Office
            continue
        relative = path.relative_to(folder).as_posix()

        try:
            kind = file_kind(path)
            if kind is None:
                report.unsupported.append(relative)
                continue
            file_hash = sha256(path)
            if file_hash in known:
                report.skipped += 1
                continue
            fields, converted = _prepare(path, kind)
        except Exception as exc:  # noqa: BLE001 — un archivo malo no frena al resto
            logger.warning("Claim %s: no se pudo registrar %s: %s", claim.pk, relative, exc)
            report.errors.append(f"{relative}: {type(exc).__name__}: {exc}")
            continue

        ClaimDocument.objects.create(
            claim=claim,
            file_path=paths.to_relative(path),
            file_name=relative,
            file_hash=file_hash,
            mime_type=MIME_BY_EXTENSION[path.suffix.lower()],
            **fields,
        )
        known.add(file_hash)
        report.registered += 1
        report.converted += converted
    return report


def _prepare(path: Path, kind: str) -> tuple[dict, bool]:
    """Campos de ClaimDocument según el tipo, y si hubo conversión."""
    if kind == "pdf":
        return pdf_fields(path), False
    if kind == "image":
        return {"page_count": 1}, False

    original = paths.to_relative(path)
    if kind == "heic":
        relative = paths.converted_relative_path(original, ".jpg")
        convert_heic_to_jpg(path, paths.resolve_data_path(relative))
        return {"page_count": 1, "converted_pdf_path": relative}, True

    relative = paths.converted_relative_path(original, ".pdf")
    target = convert_to_pdf(path, paths.resolve_data_path(relative))
    return {**pdf_fields(target), "converted_pdf_path": relative}, True


# ---------------------------------------------------------------------------
# Conversiones (se reutilizan si ya existen en _converted/)
# ---------------------------------------------------------------------------

def convert_to_pdf(source: Path, target: Path) -> Path:
    """docx, doc, rtf y xlsx -> PDF con LibreOffice (soffice --headless --convert-to pdf), guardado en target."""
    if target.exists():
        return target

    check_signature(source)
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        staged = work / f"input{OFFICE_EXTENSIONS[source.suffix.lower()]}"
        shutil.copyfile(source, staged)
        if staged.suffix == ".xlsx":
            fit_spreadsheet_for_print(staged)
        # Perfil propio por proceso: dos soffice con el mismo perfil se bloquean entre sí
        profile = f"file:///tmp/lo-profile-{os.getpid()}"
        result = subprocess.run(
            ["soffice", f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
             "--outdir", str(work / "out"), str(staged)],
            capture_output=True, text=True, timeout=SOFFICE_TIMEOUT,
        )
        produced = work / "out" / "input.pdf"
        if not produced.exists():  # soffice puede terminar con código 0 sin generar nada
            detail = (result.stderr or result.stdout or "sin salida").strip()[:300]
            raise RuntimeError(f"LibreOffice no generó el PDF: {detail}")
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), target)
    return target


_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # .doc (Word 97-2003)


def check_signature(source: Path) -> None:
    """LibreOffice abre cualquier archivo ilegible como texto plano y "convierte" basura sin error.
    Por eso se verifica que el contenido corresponda a la extensión antes de convertir."""
    ext = source.suffix.lower()
    with source.open("rb") as fh:
        head = fh.read(8)
    if ext in (".docx", ".xlsx"):
        import zipfile

        valid = zipfile.is_zipfile(source)
    elif ext == ".doc":
        valid = head == _OLE_SIGNATURE
    else:  # .rtf y .rmd
        valid = head[:5] == b"{\\rtf"
    if not valid:
        raise ValueError(f"el contenido no es un {ext} válido (archivo dañado o con otra extensión)")


MAX_COLUMN_CHARS = 60


def fit_spreadsheet_for_print(path: Path) -> None:
    """Ajusta una COPIA del Excel para que el PDF no recorte texto: LibreOffice imprime cada columna con
    su ancho original y corta lo que no cabe (un ledger perdería las descripciones de los cargos).

    Si openpyxl no puede abrirlo, se deja tal cual y LibreOffice lo convierte igual.
    """
    import openpyxl
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        logger.info("No se pudo ajustar %s para imprimir: %s", path.name, exc)
        return

    for ws in wb.worksheets:
        widths: dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                length = max(len(line) for line in str(cell.value).splitlines() or [""])
                widths[cell.column] = max(widths.get(cell.column, 0), length)
                if length > MAX_COLUMN_CHARS:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    # Sin esto, una fila con alto fijo no crece y las líneas se montan sobre la siguiente
                    ws.row_dimensions[cell.row].height = None
        for column, length in widths.items():
            ws.column_dimensions[get_column_letter(column)].width = min(length, MAX_COLUMN_CHARS) + 2
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0  # todas las páginas de alto que hagan falta
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = None
    wb.save(path)


def convert_heic_to_jpg(source: Path, target: Path) -> Path:
    if target.exists():
        return target

    from PIL import Image
    from pillow_heif import register_heif_opener

    register_heif_opener()
    target.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as img:
        img.convert("RGB").save(target, "JPEG", quality=90)
    return target


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def pdf_fields(path: Path) -> dict:
    """page_count, first_pages_text (2 primeras páginas), full_text y has_scanned_pages."""
    import pdfplumber

    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            texts.append(_clean(page.extract_text() or ""))
            page.close()  # libera la caché de la página: hay PDFs de decenas de páginas
    if not texts:
        raise ValueError("el PDF no tiene páginas")
    return {
        "page_count": len(texts),
        "first_pages_text": "\n\n".join(texts[:2]),
        "full_text": "\n\n".join(texts),
        "has_scanned_pages": any(len(t.strip()) < SCANNED_PAGE_MIN_CHARS for t in texts),
    }


def _clean(text: str) -> str:
    return text.replace("\x00", "")  # Postgres no acepta NUL en campos de texto
